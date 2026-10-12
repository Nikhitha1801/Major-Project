import io
from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "#", "Roll Number / USN", "Student Name", "Branch / Dept", "Semester",
    "Section", "Status", "Date", "Time",
]
WIDTHS = [7, 22, 28, 18, 12, 12, 14, 16, 12]


def build_attendance_workbook(
    *,
    subject_code: str,
    subject_name: str,
    faculty_name: str,
    department: str,
    section: str,
    session_date: str,
    present_students: List[Dict[str, str]],
    absent_students: List[Dict[str, str]],
) -> bytes:
    workbook = Workbook()
    report_time = datetime.now().strftime("%H:%M")
    total = len(present_students) + len(absent_students)
    attendance_pct = (len(present_students) / total * 100) if total else 0

    for status, students, title_color, summary_color, row_color in (
        ("Present", present_students, "008F68", "E8F8F1", "E8F8F1"),
        ("Absent", absent_students, "D92525", "FFF0F0", "FCE4E4"),
    ):
        sheet = workbook.active if status == "Present" else workbook.create_sheet()
        sheet.title = status
        sheet.merge_cells("A1:I1")
        sheet["A1"] = f"GSSSIETW — Attendance Report ({status})"
        sheet["A1"].fill = PatternFill("solid", fgColor="183B64")
        sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
        sheet["A1"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A2:I2")
        sheet["A2"] = f"Course: {subject_code} — {subject_name}"
        sheet["A2"].fill = PatternFill("solid", fgColor="2864DC")
        sheet["A2"].font = Font(color="FFFFFF", bold=True)
        sheet["A2"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A3:I3")
        sheet["A3"] = (
            f"Faculty: {faculty_name}   |   Dept: {department or '—'}   |   "
            f"Date: {session_date}   |   Time: {report_time}"
        )
        sheet["A3"].fill = PatternFill("solid", fgColor="DCE9FF")
        sheet["A3"].font = Font(bold=True, color="243B5A")
        sheet["A3"].alignment = Alignment(horizontal="center")

        sheet.merge_cells("A4:I4")
        sheet["A4"] = (
            f"Total {status}: {len(students)}   |   Total Enrolled: {total}   |   "
            f"Attendance: {attendance_pct:.1f}%"
        )
        sheet["A4"].fill = PatternFill("solid", fgColor=summary_color)
        sheet["A4"].font = Font(bold=True, color="243B5A")
        sheet["A4"].alignment = Alignment(horizontal="center")

        for index, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=5, column=index, value=header)
            cell.fill = PatternFill("solid", fgColor=title_color)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_number, student in enumerate(students, start=1):
            values = [
                row_number,
                student["roll_no"],
                student["name"],
                student["department"] or "—",
                student.get("semester") or "—",
                section,
                status.upper(),
                session_date,
                report_time,
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_number + 5, column=column, value=value)
                cell.fill = PatternFill("solid", fgColor=row_color)
                cell.alignment = Alignment(horizontal="center", vertical="center")

        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:I{max(5, len(students) + 5)}"
        for column, width in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in range(1, len(students) + 6):
            sheet.row_dimensions[row].height = 22

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()