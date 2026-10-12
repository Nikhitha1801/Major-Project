import io
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .. import models
from ..database import get_db
from ..core.deps import require_roles

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("/attendance/excel")
def export_attendance_excel(
    subject_id: Optional[int] = None,
    section: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_roles("admin", "hod", "faculty")),
):
    query = db.query(models.AttendanceSession)
    if subject_id:
        query = query.filter(models.AttendanceSession.subject_id == subject_id)
    if section:
        query = query.filter(models.AttendanceSession.section == section)
    if current_user.role.value == "faculty":
        query = query.filter(models.AttendanceSession.faculty_id == current_user.id)

    sessions = query.order_by(models.AttendanceSession.session_date).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    header_fill = PatternFill(start_color="4A63E8", end_color="4A63E8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Date", "Subject", "Section", "Roll No", "Student Name", "Status"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for session in sessions:
        for record in session.records:
            student_user = record.student.user
            ws.append([
                session.session_date.isoformat(),
                session.subject.name,
                session.section,
                record.student.roll_no,
                student_user.name,
                record.status.value.capitalize(),
            ])

    for col_cells in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = max_len + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "attendance_report.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
